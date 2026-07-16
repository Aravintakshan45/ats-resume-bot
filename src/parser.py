import PyPDF2
import docx
import io

def extract_text_from_pdf(file_path):
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            return text.strip()
    except Exception as e:
        raise Exception(f"Error reading PDF: {str(e)}")

def extract_text_from_docx(file_path):
    try:
        doc = docx.Document(file_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text.strip()
    except Exception as e:
        raise Exception(f"Error reading DOCX: {str(e)}")

def extract_text_from_uploaded(file_bytes, file_type):
    if file_type == 'pdf':
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()
    elif file_type == 'docx':
        doc = docx.Document(io.BytesIO(file_bytes))
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text.strip()
    else:
        raise ValueError("Unsupported file type")
